"""Claim aggregation + upsert tests, including the idempotency guarantee."""

from __future__ import annotations

import shutil
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from demurrage import config as cfg
from demurrage.claims import ClaimDraft, group_trips_to_claims, upsert_claims
from demurrage.workbook import (
    Trip,
    open_wb,
    read_assumptions,
    read_claims,
    read_deals,
    read_trips,
    save_wb,
)


# ---------- group_trips_to_claims (pure) ----------


def _trip(row: int, deal: str, consignee: str, loaded: date, demurrage_usd: float) -> Trip:
    return Trip(
        row=row,
        deal_no=deal,
        consignee=consignee,
        date_loaded=datetime(loaded.year, loaded.month, loaded.day),
        demurrage_usd=demurrage_usd,
    )


def test_group_filters_by_period_inclusive():
    trips = [
        _trip(7, "A", "C1", date(2026, 4, 30), 100),   # before
        _trip(8, "A", "C1", date(2026, 5, 1), 200),    # boundary
        _trip(9, "A", "C1", date(2026, 5, 31), 300),   # boundary
        _trip(10, "A", "C1", date(2026, 6, 1), 400),   # after
    ]
    drafts = group_trips_to_claims(trips, date(2026, 5, 1), date(2026, 5, 31))
    assert len(drafts) == 1
    assert drafts[0].amount_usd == 500
    assert drafts[0].num_trips == 2
    assert drafts[0].period_from == datetime(2026, 5, 1)
    assert drafts[0].period_to == datetime(2026, 5, 31)


def test_group_buckets_per_deal_and_emits_claim_no():
    trips = [
        _trip(7, "A", "Rubis", date(2026, 5, 1), 100),
        _trip(8, "A", "Rubis", date(2026, 5, 2), 200),
        _trip(9, "B", "Lake", date(2026, 5, 1), 50),
    ]
    drafts = group_trips_to_claims(trips, date(2026, 5, 1), date(2026, 5, 31))
    assert {d.claim_no for d in drafts} == {"A-DEM", "B-DEM"}
    a = next(d for d in drafts if d.deal_no == "A")
    assert a.amount_usd == 300
    assert a.num_trips == 2
    assert a.consignee == "Rubis"
    assert a.trip_rows == [7, 8]


def test_group_period_dates_reflect_actual_trips_not_window():
    """period_from/to come from the group's loaded dates, not the CLI window."""
    trips = [
        _trip(7, "A", "C1", date(2026, 5, 10), 100),
        _trip(8, "A", "C1", date(2026, 5, 20), 200),
    ]
    drafts = group_trips_to_claims(trips, date(2026, 5, 1), date(2026, 5, 31))
    assert drafts[0].period_from == datetime(2026, 5, 10)
    assert drafts[0].period_to == datetime(2026, 5, 20)


def test_group_skips_trips_with_no_deal_or_date():
    trips = [
        _trip(7, "A", "C1", date(2026, 5, 1), 100),
        Trip(row=8, deal_no=None, consignee="C1", date_loaded=datetime(2026, 5, 2), demurrage_usd=50),
        Trip(row=9, deal_no="B", consignee="C1", date_loaded=None, demurrage_usd=70),
    ]
    drafts = group_trips_to_claims(trips, date(2026, 5, 1), date(2026, 5, 31))
    assert len(drafts) == 1
    assert drafts[0].deal_no == "A"


def test_group_zero_demurrage_still_produces_claim():
    """A deal with eligible trips but all on-time still gets a 0-USD claim row."""
    trips = [
        _trip(7, "A", "C1", date(2026, 5, 1), 0),
        _trip(8, "A", "C1", date(2026, 5, 2), 0),
    ]
    drafts = group_trips_to_claims(trips, date(2026, 5, 1), date(2026, 5, 31))
    assert len(drafts) == 1
    assert drafts[0].amount_usd == 0
    assert drafts[0].num_trips == 2


def test_group_rejects_inverted_period():
    with pytest.raises(ValueError):
        group_trips_to_claims([], date(2026, 5, 31), date(2026, 5, 1))


# ---------- upsert_claims against a real workbook copy ----------


@pytest.fixture
def wb_copy(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    dst = tmp_path / "wb.xlsm"
    shutil.copy(cfg.WORKBOOK_PATH, dst)
    monkeypatch.setattr(cfg, "WORKBOOK_PATH", dst)
    return dst


def _claims_in_book(path: Path) -> dict[str, dict]:
    wb = load_workbook(path, keep_vba=True, data_only=False)
    try:
        out: dict[str, dict] = {}
        for c in read_claims(wb):
            if c.claim_no:
                out[c.claim_no] = {
                    "row": c.row,
                    "amount_usd": c.amount_usd,
                    "num_trips": c.num_trips,
                    "status": c.status,
                }
        return out
    finally:
        wb.close()


def test_upsert_is_idempotent(wb_copy: Path):
    """Two consecutive build_claims runs must produce the same set of claim rows
    and the same amounts. No duplicate row creation, no Sent/Paid mutation.
    """
    trips = read_trips()
    # Pick a wide window so we cover everything in the trial file.
    drafts = group_trips_to_claims(trips, date(2025, 1, 1), date(2027, 1, 1))
    assert drafts, "fixture file should produce drafts for this window"

    # Snapshot of pre-existing claim rows; we want the upsert to leave the
    # Draft rows updated in place rather than appending duplicates.
    before = _claims_in_book(wb_copy)
    issued = datetime(2026, 5, 20)

    wb = open_wb(data_only=False)
    actions1 = upsert_claims(wb, drafts, issue_date=issued)
    save_wb(wb)
    wb.close()

    after1 = _claims_in_book(wb_copy)
    # Second run with identical input.
    wb = open_wb(data_only=False)
    actions2 = upsert_claims(wb, drafts, issue_date=issued)
    save_wb(wb)
    wb.close()
    after2 = _claims_in_book(wb_copy)

    # Same claim numbers in the workbook before and after the second run.
    assert set(after2.keys()) == set(after1.keys()), \
        "second run must not introduce or remove claim numbers"
    # Same row positions — no shifting.
    for k in after2:
        assert after2[k]["row"] == after1[k]["row"], \
            f"claim {k} moved between identical runs ({after1[k]['row']} -> {after2[k]['row']})"
        assert after2[k]["amount_usd"] == after1[k]["amount_usd"]
        assert after2[k]["num_trips"] == after1[k]["num_trips"]
    # Second pass should not "create" anything.
    assert all(a != "created" for _, a in actions2), \
        f"second run created rows: {[c for c, a in actions2 if a == 'created']}"
    # And every existing draft from before is still present.
    for claim_no in before:
        assert claim_no in after2, f"existing claim {claim_no} disappeared"


def test_upsert_does_not_touch_sent_claims(wb_copy: Path):
    """Flip one existing claim to Sent, re-upsert with a different amount, verify
    Sent row is untouched.
    """
    wb = load_workbook(wb_copy, keep_vba=True, data_only=False)
    target_claim = None
    target_row = None
    for c in read_claims(wb):
        if c.status == "Draft":
            target_claim = c.claim_no
            target_row = c.row
            break
    assert target_claim, "fixture should have at least one Draft claim"

    # Lock it to "Sent" with a sentinel amount.
    ws = wb[cfg.SHEET_CLAIMS]
    ws.cell(row=target_row, column=12).value = "Sent"      # L
    ws.cell(row=target_row, column=10).value = 99999.0     # J
    save_wb(wb)
    wb.close()

    # Now try to overwrite via upsert with a different amount.
    fake = ClaimDraft(
        claim_no=target_claim,
        deal_no="ZZZ",
        consignee="Fake",
        period_from=datetime(2026, 1, 1),
        period_to=datetime(2026, 1, 31),
        num_trips=99,
        amount_usd=1.0,
        trip_rows=[1],
    )
    wb = open_wb(data_only=False)
    actions = upsert_claims(wb, [fake])
    save_wb(wb)
    wb.close()

    assert actions == [(target_claim, "skipped (Sent)")]

    reloaded = _claims_in_book(wb_copy)
    assert reloaded[target_claim]["status"] == "Sent"
    assert reloaded[target_claim]["amount_usd"] == 99999.0
    assert reloaded[target_claim]["num_trips"] != 99


def test_rerun_does_not_zero_out_amounts(wb_copy: Path):
    """Regression: openpyxl strips cached formula results on save, so a second
    read_trips() returns demurrage_usd=None. The Python-side recomputation
    (deals_by_no + assumptions) must keep amounts stable across runs.
    """
    deals_by_no = {d.deal_no: d for d in read_deals() if d.deal_no}
    assumptions = read_assumptions()

    def build_and_write():
        trips = read_trips()
        drafts = group_trips_to_claims(
            trips, date(2026, 1, 1), date(2026, 3, 31),
            deals_by_no=deals_by_no, assumptions=assumptions,
        )
        wb = open_wb(data_only=False)
        upsert_claims(wb, drafts, issue_date=datetime(2026, 5, 20))
        save_wb(wb)
        wb.close()
        return drafts

    drafts_a = build_and_write()
    drafts_b = build_and_write()
    by_no_a = {d.claim_no: d for d in drafts_a}
    by_no_b = {d.claim_no: d for d in drafts_b}
    assert set(by_no_a) == set(by_no_b)
    for k, a in by_no_a.items():
        b = by_no_b[k]
        assert a.amount_usd == b.amount_usd, (
            f"amount drift across runs for {k}: {a.amount_usd} -> {b.amount_usd}"
        )
        assert a.amount_usd > 0 or a.num_trips == 0  # at least one non-zero run


def test_upsert_appends_above_total_row(wb_copy: Path):
    """When a 'Total' row exists, new claims must land above it."""
    fake = ClaimDraft(
        claim_no="999999-DEM",
        deal_no="999999",
        consignee="Synthetic Co",
        period_from=datetime(2026, 5, 1),
        period_to=datetime(2026, 5, 31),
        num_trips=3,
        amount_usd=1234.56,
        trip_rows=[200, 201, 202],
    )
    wb = open_wb(data_only=False)
    actions = upsert_claims(wb, [fake])
    save_wb(wb)
    wb.close()

    assert actions == [("999999-DEM", "created")]
    wb = load_workbook(wb_copy, keep_vba=True, data_only=False)
    ws = wb[cfg.SHEET_CLAIMS]
    try:
        # Find the new claim and the total row, confirm ordering.
        new_row = None
        total_row = None
        for r in range(cfg.CLAIMS_FIRST_DATA_ROW, ws.max_row + 1):
            v = ws.cell(row=r, column=2).value
            if v == "999999-DEM":
                new_row = r
            elif isinstance(v, str) and v.lower() == "total":
                total_row = r
        assert new_row is not None
        assert total_row is not None
        assert new_row < total_row, f"new claim row {new_row} is at/below Total row {total_row}"
        # Total row got refreshed: amount/trips are SUM formulas spanning the
        # claims block. Excel evaluates them on next open.
        total_amount = ws.cell(row=total_row, column=10).value  # J
        total_trips = ws.cell(row=total_row, column=9).value    # I
        assert isinstance(total_amount, str) and total_amount.startswith("=SUM(J")
        assert isinstance(total_trips, str) and total_trips.startswith("=SUM(I")
        assert total_amount.endswith(f":J{total_row - 1})")
    finally:
        wb.close()
