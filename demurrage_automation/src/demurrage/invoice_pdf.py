"""Render Invoice + Annex of a single claim to PDF.

Process (PROJECT_BRIEF §3.5):
  1. Open the workbook
  2. Set Invoice!B20 = claim_no (drives all dynamic content in the template)
  3. Force recalculation
  4. Export sheets ['Invoice', 'Annex'] to PDF
  5. Save under outputs/invoices/{invoice_no}_{consignee_slug}_{YYYY-MM-DD}.pdf

Windows path uses xlwings (Excel COM, fastest, exact-match to the VBA
PublishInvoicePDF macro). Non-Windows uses LibreOffice headless.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from . import config as cfg
from .workbook import open_wb, read_claims

_RENDER_SHEETS = (cfg.SHEET_INVOICE, cfg.SHEET_ANNEX)


# ---------------------------------------------------------------------------
# Filename helpers
# ---------------------------------------------------------------------------


def _slugify_consignee(name: str) -> str:
    """'Rubis Energy Zambia' -> 'RubisEnergyZambia'. Empty -> 'Consignee'."""
    if not name:
        return "Consignee"
    cleaned = re.sub(r"[^A-Za-z0-9]+", " ", name).strip()
    return "".join(part.capitalize() for part in cleaned.split()) or "Consignee"


def _invoice_no_from_claim(claim_no: str) -> str:
    """'187296-DEM' -> 'MOC-DEM-2026-187296'. If claim_no already starts with the
    prefix, return as-is.
    """
    if claim_no.startswith(cfg.INVOICE_PREFIX):
        return claim_no
    deal_part = claim_no.split("-", 1)[0]
    return f"{cfg.INVOICE_PREFIX}{deal_part}"


def invoice_pdf_path(
    claim_no: str,
    consignee: str,
    *,
    issue_date: Optional[date] = None,
    output_dir: Optional[Path] = None,
) -> Path:
    """Compose the output PDF path (does not create the file)."""
    d = issue_date or date.today()
    out_dir = output_dir or cfg.INVOICES_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    invoice_no = _invoice_no_from_claim(claim_no)
    slug = _slugify_consignee(consignee)
    return out_dir / f"{invoice_no}_{slug}_{d:%Y-%m-%d}.pdf"


# ---------------------------------------------------------------------------
# Lookup
# ---------------------------------------------------------------------------


def _lookup_consignee(claim_no: str) -> str:
    """Find the consignee for `claim_no` so the filename slug is accurate."""
    for c in read_claims():
        if c.claim_no == claim_no and c.consignee:
            return c.consignee
    return ""


# ---------------------------------------------------------------------------
# Public entrypoint
# ---------------------------------------------------------------------------


def render_invoice_pdf(
    claim_no: str,
    *,
    source_path: Optional[Path] = None,
    output_dir: Optional[Path] = None,
    issue_date: Optional[date] = None,
    consignee: Optional[str] = None,
) -> Path:
    """Render one claim's Invoice+Annex to PDF. Returns the path written."""
    source = source_path or cfg.WORKBOOK_PATH
    if not source.exists():
        raise FileNotFoundError(f"Source workbook not found: {source}")

    consignee = consignee if consignee is not None else _lookup_consignee(claim_no)
    out_path = invoice_pdf_path(
        claim_no, consignee, issue_date=issue_date, output_dir=output_dir
    )

    if sys.platform == "win32":
        _render_via_xlwings(claim_no, source, out_path)
    else:
        _render_via_libreoffice(claim_no, source, out_path)

    if not out_path.exists() or out_path.stat().st_size == 0:
        raise RuntimeError(f"PDF was not produced at {out_path}")
    return out_path


# ---------------------------------------------------------------------------
# Windows: xlwings + Excel COM
# ---------------------------------------------------------------------------


def _render_via_xlwings(claim_no: str, source: Path, out_path: Path) -> None:
    """Drive Excel via xlwings. Matches the VBA PublishInvoicePDF semantics."""
    import xlwings as xw  # imported lazily so non-Windows installs don't need it

    app = xw.App(visible=False, add_book=False)
    try:
        app.display_alerts = False
        app.screen_updating = False
        book = app.books.open(str(source.resolve()), update_links=False)
        try:
            book.sheets[cfg.SHEET_INVOICE].range("B20").value = claim_no
            app.calculate()
            sheets = book.api.Sheets([cfg.SHEET_INVOICE, cfg.SHEET_ANNEX])
            sheets.Select()
            book.api.ActiveSheet.ExportAsFixedFormat(
                Type=0,  # xlTypePDF
                Filename=str(out_path.resolve()),
                Quality=0,  # xlQualityStandard
                IgnorePrintAreas=False,
                OpenAfterPublish=False,
            )
            book.sheets[cfg.SHEET_INVOICE].select()
        finally:
            book.close()  # discard the B20 mutation
    finally:
        app.quit()


# ---------------------------------------------------------------------------
# Linux/macOS: LibreOffice headless
# ---------------------------------------------------------------------------


def _libreoffice_binary() -> str:
    for name in ("libreoffice", "soffice"):
        path = shutil.which(name)
        if path:
            return path
    raise RuntimeError(
        "LibreOffice not found on PATH. Install libreoffice or run on Windows."
    )


def _render_via_libreoffice(claim_no: str, source: Path, out_path: Path) -> None:
    """Headless render:
      1. copy source to a temp xlsm
      2. write claim_no into Invoice!B20
      3. hide every sheet except Invoice and Annex (so soffice exports only those)
      4. save
      5. soffice --headless --convert-to pdf
      6. move result to out_path
    """
    binary = _libreoffice_binary()
    with tempfile.TemporaryDirectory(prefix="demurrage_pdf_") as tmpdir:
        tmp = Path(tmpdir)
        staged = tmp / source.name
        shutil.copy(source, staged)

        wb = load_workbook(staged, keep_vba=True, data_only=False)
        try:
            inv = wb[cfg.SHEET_INVOICE]
            inv["B20"].value = claim_no

            # LibreOffice respects sheet visibility; hide every other sheet so
            # the resulting PDF contains only Invoice + Annex pages, same as
            # the VBA PublishInvoicePDF macro.
            kept = set(_RENDER_SHEETS)
            for name in wb.sheetnames:
                if name in kept:
                    wb[name].sheet_state = "visible"
                else:
                    wb[name].sheet_state = "hidden"
            # At least one sheet must remain visible — guard against typos.
            if not any(wb[s].sheet_state == "visible" for s in wb.sheetnames):
                raise RuntimeError("No visible sheets after filtering")
            wb.active = wb.sheetnames.index(cfg.SHEET_INVOICE)

            wb.save(staged)
        finally:
            wb.close()

        # LibreOffice ignores SOFFICE userprofile contention if we give it a
        # private one per call — important for parallel runs (M6 dunning sweep).
        profile_dir = tmp / "profile"
        env_args = [f"-env:UserInstallation=file://{profile_dir}"]

        proc = subprocess.run(
            [
                binary,
                *env_args,
                "--headless",
                "--nologo",
                "--nofirststartwizard",
                "--norestore",
                "--convert-to",
                "pdf",
                "--outdir",
                str(tmp),
                str(staged),
            ],
            capture_output=True,
            text=True,
            timeout=180,
        )
        if proc.returncode != 0:
            raise RuntimeError(
                f"LibreOffice exited {proc.returncode}: "
                f"{(proc.stderr or proc.stdout).strip()}"
            )

        produced = staged.with_suffix(".pdf")
        if not produced.exists():
            # soffice writes to <outdir>/<input-stem>.pdf
            candidates = list(tmp.glob("*.pdf"))
            if not candidates:
                raise RuntimeError(
                    f"LibreOffice produced no PDF. stdout={proc.stdout!r} stderr={proc.stderr!r}"
                )
            produced = candidates[0]

        out_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(produced), str(out_path))


# ---------------------------------------------------------------------------
# Batch helper
# ---------------------------------------------------------------------------


def render_for_claim_no(claim_no: str, **kwargs) -> Path:
    """Convenience wrapper; identical to render_invoice_pdf."""
    return render_invoice_pdf(claim_no, **kwargs)
